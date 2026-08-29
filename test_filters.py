#!/usr/bin/env python3
"""Regression suite for jobhunt.py - plain stdlib, no pytest.

Run:  python test_filters.py
Covers: is_us_location, intern detection, visa signal, skill word-bounding,
HTML cleaning, and a write_excel smoke test against the sheet contract.
"""
import os
import sys
import tempfile
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import jobhunt  # noqa: E402

PASS = 0
FAIL = 0
FAILURES = []


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"PASS  {name}")
    else:
        FAIL += 1
        FAILURES.append(name)
        print(f"FAIL  {name}" + (f"  ({detail})" if detail else ""))


# ---------------------------------------------------------------- 1. US filter
def test_us_locations():
    false_cases = ["Remote UK", "Remote - Europe", "Cambridge, UK",
                   "San Jose, Costa Rica", "London", "Toronto, ON",
                   "Remote EMEA", "Bengaluru", "Paris, France", "Sydney"]
    true_cases = ["New York, NY", "Remote - US", "Springfield, MO",
                  "Boise, ID", "United States", "SF Bay Area",
                  "San Jose, CA", "Austin, Texas", "Remote (US)"]
    for loc in false_cases:
        got = jobhunt.is_us_location(loc)
        check(f"is_us_location({loc!r}) is False", got is False, f"got {got!r}")
    for loc in true_cases:
        got = jobhunt.is_us_location(loc)
        check(f"is_us_location({loc!r}) is True", got is True, f"got {got!r}")
    got = jobhunt.is_us_location("London | New York, NY")
    check("multi-location 'London | New York, NY' -> True", got is True,
          f"got {got!r}")
    got = jobhunt.is_us_location("London | Paris")
    check("multi-location 'London | Paris' -> False", got is False,
          f"got {got!r}")


# ------------------------------------------------------------ 2. intern titles
CFG = {
    "target_titles": ["product", "pm", "manager"],
    "secondary_titles": [],
    "locations": [],
    "skills": [],
    "exclude_title_words": [],
    "us_only": True,
    "strict_location": False,
    "drop_no_sponsorship": False,
    "want_internships": True,
    "want_fulltime": True,
}


def score(title, desc="", loc="New York, NY", cfg=None):
    job = {"title": title, "location": loc, "description": desc,
           "url": "https://example.com/j"}
    return jobhunt.score_job(job, cfg or CFG)


def test_intern_detection():
    cases = [
        ("Product Manager | International", False),
        ("Product Management Intern", True),
        ("Internal Tools PM", False),
        ("Internet Services Manager", False),
    ]
    for title, want_intern in cases:
        res = score(title)
        check(f"score_job kept {title!r}", res is not None)
        if res is None:
            continue
        level = res[3]
        if want_intern:
            check(f"{title!r} classified Internship", level == "Internship",
                  f"level={level!r}")
        else:
            check(f"{title!r} NOT classified Internship",
                  level != "Internship", f"level={level!r}")


# -------------------------------------------------------------- 3. visa signal
def test_visa():
    neg = jobhunt.classify_visa_signal(
        "We are unable to provide H-1B sponsorship.")
    check("'unable to provide H-1B sponsorship' -> NO_SPONSORSHIP",
          neg == "NO_SPONSORSHIP", f"got {neg!r}")
    pos = jobhunt.classify_visa_signal("We sponsor H-1B visas.")
    check("'We sponsor H-1B visas.' -> SPONSORS",
          pos == "SPONSORS", f"got {pos!r}")

    # And through score_job: the negative posting must not carry a positive flag
    cfg = dict(CFG)
    res = score("Product Manager",
                desc="We are unable to provide H-1B sponsorship.", cfg=cfg)
    if res is not None:
        visa_flag = res[2]
        check("score_job visa flag on negative posting is not positive",
              not visa_flag.startswith("SPONSORS"), f"flag={visa_flag!r}")
    else:
        # drop_no_sponsorship is False in CFG, so it should not be dropped
        check("score_job visa flag on negative posting is not positive",
              False, "row was dropped despite drop_no_sponsorship=False")
    res = score("Product Manager", desc="We sponsor H-1B visas.", cfg=cfg)
    check("score_job flags sponsoring posting positive",
          res is not None and res[2].startswith("SPONSORS"),
          f"res={res!r}")


# ------------------------------------------------------------- 4. skill bounds
def test_skills():
    cfg = dict(CFG)
    cfg["skills"] = ["Java", "RAG"]
    res = score("Product Manager",
                desc="We use JavaScript to leverage our platform", cfg=cfg)
    check("JavaScript/leverage matches neither Java nor RAG",
          res is not None and "your skills" not in res[1], f"res={res!r}")
    res = score("Product Manager", desc="Java and RAG pipelines", cfg=cfg)
    ok = (res is not None and "your skills" in res[1]
          and "Java" in res[1] and "RAG" in res[1])
    check("'Java and RAG pipelines' matches both", ok, f"res={res!r}")


# ------------------------------------------------------ 5. HTML strip/unescape
def test_clean_html():
    out = jobhunt.clean_html("&lt;p&gt;Great &amp;amp; fun role&lt;/p&gt;")
    check("clean_html flattens escaped Greenhouse HTML",
          "<" not in out and "&amp;" not in out and "Great & fun role" in out,
          f"got {out!r}")


# ------------------------------------------------------- 6. write_excel smoke
def make_row(key, company, ai_fit, elig, days_old, title="Product Manager"):
    r = {c: "" for c in jobhunt.COLUMNS}
    r.update({
        "AI Fit": ai_fit, "Score": 60, "Status": "To Review",
        "Company": company, "Title": title, "Level": "Mid-level",
        "Location": "New York, NY", "Posted": "2026-08-28",
        "Days Old": days_old, "Eligibility": elig,
        "Apply Link": f"https://example.com/{key}", "Key": key,
    })
    return r


def test_write_excel():
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    fresh_row = make_row("k1", "FreshCo", 90, "YES", 1)
    no_row = make_row("k2", "BlockedCo", 80, "NO", 5)
    unscored = make_row("k3", "MysteryCo", "", "", 10)
    rows = [fresh_row, no_row, unscored]
    sources = [
        {"name": "TestCo", "ats": "greenhouse", "status": "OK",
         "count": 10, "kept": 2, "error": ""},
        {"name": "Simplify New-Grad", "ats": "github", "status": "FAILED",
         "count": 0, "kept": 0, "error": "HTTP 500"},
    ]
    tmpdir = tempfile.mkdtemp(prefix="jobhunt_test_")
    path = os.path.join(tmpdir, "test_out.xlsx")
    jobhunt.write_excel(rows, list(rows), {"k1"}, [], path,
                        lowfit=[], sources=sources, incomplete=True)

    wb = load_workbook(path)
    expect_order = ["Dashboard", "Fresh (48h)", "Ready to Apply",
                    "New This Run", "All Matches", "Low Fit", "Run Report",
                    "Quarantine (dead links)"]
    check("sheet order matches contract", wb.sheetnames == expect_order,
          f"got {wb.sheetnames}")

    def companies(sheet):
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        if "Company" not in headers:
            return []
        i = headers.index("Company")
        return [row[i] for row in ws.iter_rows(min_row=2, values_only=True)
                if row[i]]

    ready_cos = companies("Ready to Apply")
    low_cos = companies("Low Fit")
    fresh_cos = companies("Fresh (48h)")
    check("Eligibility NO row is on Low Fit", "BlockedCo" in low_cos,
          f"Low Fit={low_cos}")
    check("Eligibility NO row is NOT on Ready to Apply",
          "BlockedCo" not in ready_cos, f"Ready={ready_cos}")
    check("unscored row IS on Ready to Apply", "MysteryCo" in ready_cos,
          f"Ready={ready_cos}")
    check("unscored row is NOT on Low Fit", "MysteryCo" not in low_cos,
          f"Low Fit={low_cos}")
    check("fresh row is on the Fresh sheet", "FreshCo" in fresh_cos,
          f"Fresh={fresh_cos}")
    check("stale rows are NOT on the Fresh sheet",
          "MysteryCo" not in fresh_cos and "BlockedCo" not in fresh_cos,
          f"Fresh={fresh_cos}")

    # Run Report: source rows present
    ws_rep = wb["Run Report"]
    cells = [str(c.value) for row in ws_rep.iter_rows() for c in row
             if c.value not in (None, "")]
    check("Run Report lists TestCo source",
          any("TestCo" == v for v in cells), f"cells={cells[:20]}")
    check("Run Report lists the Simplify feed",
          any("Simplify New-Grad" == v for v in cells))
    check("Run Report carries the INCOMPLETE summary",
          any("INCOMPLETE" in v for v in cells))
    check("Run Report shows the error text",
          any("HTTP 500" in v for v in cells))

    # Status dropdown targets the Status column, located BY HEADER
    ws = wb["Ready to Apply"]
    headers = [c.value for c in ws[1]]
    status_letter = get_column_letter(headers.index("Status") + 1)
    dv_ok = False
    for dv in ws.data_validations.dataValidation:
        if dv.formula1 and "Applied" in dv.formula1:
            for rng in dv.sqref.ranges:
                if str(rng).startswith(f"{status_letter}2"):
                    dv_ok = True
    check("Status dropdown targets the Status column", dv_ok,
          f"Status col={status_letter}, dvs="
          f"{[(d.formula1, str(d.sqref)) for d in ws.data_validations.dataValidation]}")

    try:
        os.remove(path)
        os.rmdir(tmpdir)
    except OSError:
        pass


def main():
    global FAIL
    tests = [test_us_locations, test_intern_detection, test_visa,
             test_skills, test_clean_html, test_write_excel]
    for t in tests:
        print(f"--- {t.__name__} ---")
        try:
            t()
        except Exception:
            FAIL += 1
            FAILURES.append(f"{t.__name__} (exception)")
            print(f"FAIL  {t.__name__} raised:")
            traceback.print_exc()
    print()
    print(f"{PASS} passed, {FAIL} failed")
    if FAILURES:
        print("Failed cases:")
        for f in FAILURES:
            print(f"  - {f}")
        sys.exit(1)
    sys.exit(0)


if __name__ == "__main__":
    main()
