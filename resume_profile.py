"""
Resume -> profile object.

`jobhunt.py` and `ai_rank.py` were written around one hard-coded candidate:
ai_rank.CANDIDATE is a literal briefing about Shruti, and config.json carries her
titles and locations. This module removes that assumption. It reads whatever
resume the user uploads, and asks Gemini to write the *equivalent* briefing for
that person - including a 0-100 rubric derived from their own timeline rather
than from Shruti's.

Public API (exactly two functions):

    extract_text(path)              -> str
    build_profile(path, today=None) -> dict   # raises ValueError on failure

The dict it returns is the profile object from the HTTP contract; the web layer
writes it to profile.json and ai_rank.py reads candidate_block out of it in
place of its own CANDIDATE constant.

The single most important thing this module gets right is ELIGIBILITY. A rubric
that puts "Summer 2027 PM internship" at 90-100 is correct for a first-semester
master's student and actively harmful for someone who graduated four years ago.
So the prompt below refuses to hard-code any band: it hands the model today's
date and makes it reason about the newest education end date first, then build
the bands from that conclusion.

Two post-audit guarantees:
- PII redaction: the Gemini prompt is built from a redacted COPY of the resume
  text - email addresses -> "[email]", phone numbers -> "[phone]",
  linkedin/github/portfolio URLs -> "[link]", street addresses -> "[address]".
  The candidate's name is kept (the profile needs it), and the output of
  extract_text is never altered - only the prompt copy is.
- No invented preferences: the prompt forbids locations the resume does not
  evidence ("locations" is resume-evidenced cities/stated preferences plus
  "Remote", or exactly ["Remote"] if none), and visa status is never asserted
  as fact - at most "likely requires sponsorship (inferred from education
  history — confirm with the candidate)".

No new dependencies: PDF via fitz (PyMuPDF), DOCX via python-docx, and the
Gemini call is borrowed wholesale from ai_rank (_api_key + _call) so there is
exactly one HTTP client in this project.
"""

import datetime
import json
import os
import re

import ai_rank

# Resume text handed to the model. Two dense pages is ~6k chars; 20k leaves
# room for a long CV without paying for a novel.
MAX_RESUME_CHARS = 20000

# A resume that extracts to less than this is almost certainly a scanned image.
MIN_RESUME_CHARS = 200

SUPPORTED_EXTS = (".pdf", ".docx", ".txt")

# Anything shorter than this is not a usable briefing, whatever the model says.
MIN_BLOCK_CHARS = 200
MAX_BLOCK_CHARS = 12000

_LIST_KEYS = ("target_titles", "secondary_titles", "exclude_title_words",
              "skills", "locations")

# Last-resort list values, used only when the model omits a key AND config.json
# has nothing to offer. Deliberately generic - the real values come from the
# resume.
_HARD_DEFAULTS = {
    "target_titles": ["Product Manager"],
    "secondary_titles": ["Business Analyst", "Program Manager"],
    "exclude_title_words": ["sales", "account executive", "recruiter"],
    "skills": ["SQL", "analytics", "stakeholder management"],
    "locations": ["United States", "Remote"],
}

# Zero-width and BOM characters that PDF extraction likes to smuggle in.
_INVISIBLE = re.compile("[\u200b\u200c\u200d\u2060\ufeff\xa0]")


# --------------------------------------------------------------------------
# PII redaction (applied to a COPY used for the Gemini prompt only)
# --------------------------------------------------------------------------

_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")

# Anything URL-shaped. A callback decides whether it is a personal link
# (linkedin/github/portfolio) worth masking; other tokens that merely look
# like URLs ("Node.js", a company site in an employer line) are left alone.
_URLISH_RE = re.compile(
    r"(?:https?://|www\.)[^\s<>()\[\]|]+"
    r"|\b[A-Za-z0-9][A-Za-z0-9.-]*\.[A-Za-z]{2,}\b(?:/[^\s<>()\[\]|]*)?",
    re.I)
_PII_HOST = re.compile(r"linkedin|github|portfolio", re.I)

# Digit runs (7-16 digits) with common phone separators. \s is deliberately
# NOT in the class - a phone never spans lines, but "2023\n2024" would.
_PHONE_RE = re.compile(r"\(?\+?\d[\d \t().-]{5,}\d")
_YEAR_RANGE = re.compile(r"(?:19|20)\d\d\s*[-\u2013\u2014/]?\s*(?:19|20)\d\d")

# House number + Title-Case street name + a street suffix, optional apt/unit.
# Title-case-only on the name words keeps "shipped 5 product road maps" safe.
# Horizontal whitespace ONLY ([ \t]) - like a phone, a street address never
# spans lines. With \s+ here, "NY 10027\nShruti Gupta\nStreet-smart leader"
# redacted the candidate's NAME into "[address]", and a match could start at a
# digit on the previous line and swallow that line's real content too.
_ADDRESS_RE = re.compile(
    r"\b\d{1,6}(?:[ \t]+[A-Z0-9][A-Za-z0-9.'-]*){1,4}[ \t]+"
    r"(?i:street|st|avenue|ave|road|rd|lane|ln|drive|dr|boulevard|blvd"
    r"|court|ct|place|pl|terrace|ter|way|circle|cir)\b\.?"
    r"(?i:,?[ \t]*(?:apt|apartment|suite|ste|unit|#)\.?[ \t]*[\w-]+)?")


def _phone_sub(m):
    s = m.group(0)
    digits = re.sub(r"\D", "", s)
    if not 7 <= len(digits) <= 16:
        return s          # too few/many digits to be a phone number
    if _YEAR_RANGE.fullmatch(s.strip("() ")):
        return s          # "2019 - 2023" is a tenure, not a phone
    return "[phone]"


def _redact_pii(text):
    """Return a copy of `text` with contact PII masked, for the prompt ONLY.

    Emails -> [email], linkedin/github/portfolio URLs -> [link], street
    addresses -> [address], phone numbers -> [phone]. The candidate's name is
    deliberately kept - the profile needs it. Callers must never write this
    copy back anywhere; extract_text output stays exactly as extracted.
    """
    text = _EMAIL_RE.sub("[email]", text)
    text = _URLISH_RE.sub(
        lambda m: "[link]" if _PII_HOST.search(m.group(0)) else m.group(0),
        text)
    text = _ADDRESS_RE.sub("[address]", text)
    text = _PHONE_RE.sub(_phone_sub, text)
    return text


# --------------------------------------------------------------------------
# text extraction
# --------------------------------------------------------------------------

def _normalize(text):
    """Collapse horizontal whitespace but keep the line structure.

    A resume's meaning lives in its line breaks - "Columbia University" and
    "Aug 2026 - Dec 2027" are two separate facts on two separate lines, and
    flattening them into one paragraph is how you end up with a model that
    attaches the graduation date to the previous employer.
    """
    text = _INVISIBLE.sub(" ", text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t\f\v]+", " ", ln).strip() for ln in text.split("\n")]
    text = "\n".join(lines)
    # Three or more blank lines carry no information; one does.
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_text(path):
    try:
        import fitz  # PyMuPDF
    except ImportError as e:  # pragma: no cover - fitz ships with this project
        raise ValueError("PyMuPDF (fitz) is not installed, so PDF resumes "
                         "cannot be read. Upload a .docx or .txt instead.") from e

    try:
        doc = fitz.open(path)
    except Exception as e:
        raise ValueError("Could not open that PDF (%s). It may be corrupt or "
                         "not really a PDF." % e) from e

    try:
        # needs_pass stays True until a password is supplied; is_encrypted can
        # remain True even on a document we are allowed to read, so needs_pass
        # is the flag that actually blocks us.
        if getattr(doc, "needs_pass", False):
            raise ValueError("That PDF is password-protected. Remove the "
                             "password (or export an unlocked copy) and upload "
                             "it again.")
        pages = []
        for page in doc:
            try:
                pages.append(page.get_text())
            except Exception:
                continue
        if not pages:
            raise ValueError("That PDF has no readable pages.")
        return "\n".join(pages)
    finally:
        try:
            doc.close()
        except Exception:
            pass


def _docx_tables(tables, out, seen):
    """Walk tables (and nested tables) collecting cell text.

    Plenty of resume templates are one borderless 2-column table: dates on the
    right, everything else on the left. Reading only doc.paragraphs on one of
    those returns an almost empty string.
    """
    for table in tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                # Horizontally merged cells repeat the same _tc element across
                # the row; id() de-dupes them so text is not tripled.
                cid = id(cell._tc)
                if cid in seen:
                    continue
                seen.add(cid)
                txt = "\n".join(p.text for p in cell.paragraphs).strip()
                if txt:
                    cells.append(txt)
                if cell.tables:
                    _docx_tables(cell.tables, out, seen)
            if cells:
                out.append(" | ".join(cells))


def _docx_text(path):
    try:
        from docx import Document
    except ImportError as e:  # pragma: no cover - python-docx ships here
        raise ValueError("python-docx is not installed, so .docx resumes "
                         "cannot be read. Upload a PDF or .txt instead.") from e

    try:
        document = Document(path)
    except Exception as e:
        raise ValueError("Could not open that Word file (%s). If it is an old "
                         ".doc, re-save it as .docx or PDF." % e) from e

    parts = [p.text for p in document.paragraphs]
    _docx_tables(document.tables, parts, set())

    for section in document.sections:
        for container in (section.header, section.footer):
            try:
                parts.extend(p.text for p in container.paragraphs)
            except Exception:
                continue

    return "\n".join(p for p in parts if p and p.strip())


def _txt_text(path):
    with open(path, "rb") as f:
        raw = f.read()
    # utf-8-sig first: a .txt saved by Notepad, or produced by a PowerShell
    # redirect, carries a BOM. cp1252 second: "ANSI" is still the save default
    # in plenty of Windows editors, and decoding one of those as utf-8 with
    # errors="replace" turns every curly quote, en dash and bullet into U+FFFD
    # before the model ever sees the resume.
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def extract_text(path):
    """Return the plain text of a .pdf / .docx / .txt resume.

    Raises ValueError with a message fit to show a user: unsupported type,
    encrypted PDF, or too little text to work with.
    """
    path = os.fspath(path)
    if not os.path.isfile(path):
        raise ValueError("No such file: %s" % path)

    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        raw = _pdf_text(path)
    elif ext == ".docx":
        raw = _docx_text(path)
    elif ext in (".txt", ".text", ".md"):
        raw = _txt_text(path)
    elif ext == ".doc":
        raise ValueError("Old .doc files are not supported. Open it in Word, "
                         "save as .docx or PDF, and upload again.")
    else:
        raise ValueError("Unsupported file type '%s'. Upload one of: %s."
                         % (ext or path, ", ".join(SUPPORTED_EXTS)))

    text = _normalize(raw)
    if len(text) < MIN_RESUME_CHARS:
        raise ValueError(
            "Only %d characters of text came out of that file. If it is a PDF "
            "it is probably a scan or an image export with no text layer - "
            "export a text-based PDF from Word or Google Docs, or upload the "
            ".docx instead." % len(text))
    return text


# --------------------------------------------------------------------------
# the prompt
# --------------------------------------------------------------------------

# This skeleton mirrors ai_rank.CANDIDATE section for section (CURRENT SITUATION
# / BACKGROUND / WHAT THEY WANT / HOW TO SCORE FIT (0-100) with six bands /
# BOOST / PENALISE) so the generated block is a drop-in replacement for that
# constant. Placeholders only - never a worked example, or the model copies the
# example candidate's facts into someone else's profile.
_BLOCK_SKELETON = """\
<FULL NAME> - target profile

CURRENT SITUATION (this is the most important part):
- <exactly where they are today: degree in progress and its graduation date, or
  graduated and currently in <role> with <N> years of experience>
- <work authorisation: only what the resume states outright; if the education
  timeline merely implies it, hedge - "likely requires sponsorship (inferred
  from education history — confirm with the candidate)" - never assert visa
  status as fact>
- <where they are based per the resume, plus Remote - never a city the resume
  does not mention>

BACKGROUND:
- <career arc: employers, titles, dates, the level they reached>
- <domain depth: the problem spaces they have actually shipped in>
- <technical stack and tooling>
- <education, degrees, GPA if the resume states one>

WHAT THEY WANT: <primary function>. Secondarily <adjacent functions that lead
back to the primary one>.

HOW TO SCORE FIT (0-100):
- 90-100: <the sweet spot, straight out of the timeline reasoning>
- 75-89: <close variants of the sweet spot>
- 55-74: <plausible but imperfect: right level and weaker domain, or right
  domain and slightly wrong level>
- 30-54: <worth tracking, not worth applying to now - say why>
- 10-29: <wrong level, or a weak domain match>
- 0-9: <hard disqualifiers: seniority far beyond them, wrong function,
  no-sponsorship or clearance-only postings, roles they cannot start>

BOOST for: <the domains, industries and keywords they have genuinely shipped and
could defend in an interview>.

PENALISE: <domains and role types that would waste their time>.
"""

_PROMPT = """\
You are a career-intelligence analyst. Read the RESUME at the bottom and return
ONE JSON object describing this candidate. An automated job search uses it to
score live US job postings from 0 to 100, so every word has to be actionable.

Contact details were masked before you received this resume: "[email]",
"[phone]", "[link]" and "[address]" are redaction markers, not content. Never
copy a marker into any field and never try to reconstruct what one hid.

TODAY'S DATE IS @@TODAY@@. Work out the candidate's timeline before you write
anything, and let that conclusion drive the entire rubric:

1. Find the education entry with the newest end date.
2. If that end date is AFTER @@TODAY@@, this person is a CURRENT STUDENT. They
   cannot start a full-time job before that date. So internships / co-ops for
   the next available season, and new-grad or rotational programs starting at or
   after graduation, belong in the top bands - and ordinary full-time mid-level
   roles belong LOW, not because the work is a bad match but because they are
   not available to start. Name the specific season and year (for example
   "Summer 2027") in the bands.
3. If that end date is BEFORE @@TODAY@@, this person has GRADUATED. Add up their
   full-time work history through @@TODAY@@ to get N years of experience, and
   state N explicitly. Full-time roles at that level are the top band. Student
   internships and new-grad programs score near zero because they are no longer
   eligible for them. Senior / Staff / Principal roles score low unless N
   actually supports them.
4. Work out which rung of the ladder they are on and what the next rung is
   called at a normal US company.
5. Work authorisation, honestly: never state visa status as a fact unless the
   resume itself states it. If the education timeline merely implies it (for
   example a current international student), the candidate_block must hedge -
   write "likely requires sponsorship (inferred from education history —
   confirm with the candidate)" - never "is on F-1" or "needs an H-1B" as
   fact. When sponsorship is likely needed, postings that refuse
   sponsorship/OPT/CPT or demand citizenship or a security clearance are 0-9.
6. "locations" must contain ONLY locations the resume evidences: the city they
   currently live, work or study in, and any location preference the resume
   states outright - plus "Remote". Do NOT add cities the resume never
   mentions (no guessing job-market hubs like San Francisco or Seattle). If
   the resume evidences no location at all, return exactly ["Remote"].
7. Never invent an employer, a school, a date, a skill or a location that is
   not in the resume. If something is genuinely ambiguous, pick the reading the
   resume best supports and stay consistent with it.

Return ONLY this JSON object - no markdown fences, no commentary:

{
  "name": "the candidate's full name exactly as written on the resume",
  "headline": "one line, at most 90 characters, e.g. 'Product Manager - ex-Senior SWE, Columbia MS 2027'",
  "summary": "2-3 plain sentences: who they are, what they have shipped, what they are looking for now. Written for the candidate to read about themselves.",
  "target_titles": ["8-14 exact job titles to search for - the roles they should get hired into next"],
  "secondary_titles": ["6-12 adjacent titles worth surfacing but ranked below the targets"],
  "exclude_title_words": ["8-15 lowercase words or phrases that mean a posting is NOT for them - wrong function (e.g. 'account executive'), or wrong seniority given their timeline"],
  "skills": ["12-25 concrete skills, tools and domains taken from the resume"],
  "locations": ["ONLY locations evidenced in the resume (current city, explicitly stated preferences), most preferred first, plus 'Remote'; exactly ['Remote'] if the resume evidences none"],
  "candidate_block": "the briefing described below, as one JSON string with \\n line breaks"
}

THE candidate_block IS THE MOST IMPORTANT FIELD. It is pasted verbatim into the
prompt that scores every posting, so it has to stand alone: someone who has
never seen the resume should be able to score jobs correctly from it and nothing
else. Follow this structure exactly - same sections, same order, same six
scoring bands. Replace every <placeholder> with this candidate's real facts and
delete the angle brackets. Do not copy the placeholder wording.

@@SKELETON@@

Rules for candidate_block:
- 250 to 600 words. Plain text with real line breaks (escaped as \\n in JSON).
- Every band must name concrete role types, and must mention dates, seasons or
  years of experience wherever eligibility depends on them.
- The bands must be mutually exclusive and must cover the whole 0-100 range.
- The 90-100 band has to be reachable: it describes jobs that exist and that
  this person could actually start.
- Write it in the third person, using the candidate's name or "they".

RESUME:
---
@@RESUME@@
---
"""


def _build_prompt(resume_text, today):
    return (_PROMPT
            .replace("@@TODAY@@", today)
            .replace("@@SKELETON@@", _BLOCK_SKELETON)
            .replace("@@RESUME@@", resume_text))


# --------------------------------------------------------------------------
# validation
# --------------------------------------------------------------------------

def _parse_object(text):
    """Pull a JSON object out of the model's reply."""
    text = (text or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-z]*\n?|```$", "", text, flags=re.M).strip()
    data = None
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", text, re.S)
        if m:
            try:
                data = json.loads(m.group(0))
            except json.JSONDecodeError:
                data = None
    if isinstance(data, list):  # occasionally wrapped in an array
        data = next((d for d in data if isinstance(d, dict)), None)
    if not isinstance(data, dict):
        raise ValueError("The AI did not return a usable profile (its reply "
                         "was not JSON). Try uploading the resume again.")
    return data


def _clean_str(value, limit, default=""):
    if not isinstance(value, str):
        return default
    value = _INVISIBLE.sub(" ", value).strip()
    value = re.sub(r"[ \t]+", " ", value)
    if not value:
        return default
    return value[:limit].strip()


def _clean_list(value, default, item_limit=80, max_items=40, lower=False):
    """Coerce whatever came back into a de-duplicated list of short strings."""
    if isinstance(value, str):
        value = re.split(r"[;,\n]", value)
    if not isinstance(value, (list, tuple)):
        return list(default)

    out, seen = [], set()
    for item in value:
        if isinstance(item, (int, float)) and not isinstance(item, bool):
            item = str(item)
        if not isinstance(item, str):
            continue
        item = re.sub(r"\s+", " ", _INVISIBLE.sub(" ", item)).strip(" -•\t")
        if not item or len(item) > item_limit:
            continue
        if lower:
            item = item.lower()
        key = item.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(item)
        if len(out) >= max_items:
            break
    return out or list(default)


def _config_defaults():
    """Fall back to whatever config.json already holds before inventing values.

    If the model drops a list, the search config the user is already running is
    a much better guess than anything hard-coded in this file.
    """
    defaults = dict(_HARD_DEFAULTS)
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")
    try:
        # utf-8-sig, matching ai_rank: a config.json the user has opened in
        # Notepad comes back with a BOM, json.load raises on it, and every
        # fallback list would silently collapse to the generic hard defaults.
        with open(path, encoding="utf-8-sig") as f:
            cfg = json.load(f)
    except (OSError, ValueError):
        return defaults
    if not isinstance(cfg, dict):
        return defaults
    for key in _LIST_KEYS:
        value = cfg.get(key)
        if isinstance(value, list):
            items = [v for v in value if isinstance(v, str) and v.strip()]
            if items:
                defaults[key] = items[:40]
    return defaults


def _validate(data):
    """Turn the model's dict into the profile object, or raise ValueError."""
    defaults = _config_defaults()

    block = data.get("candidate_block")
    if isinstance(block, (list, tuple)):
        block = "\n".join(str(x) for x in block)
    if not isinstance(block, str):
        block = ""
    block = _INVISIBLE.sub(" ", block).replace("\r\n", "\n").replace("\r", "\n")
    block = "\n".join(ln.rstrip() for ln in block.split("\n")).strip()

    # candidate_block is the whole point of the call. A stub is a failure, not
    # something to paper over with a default - a bad rubric silently ruins every
    # score in the run.
    if len(block) < MIN_BLOCK_CHARS:
        raise ValueError("The AI returned an incomplete profile - the scoring "
                         "briefing was empty or far too short. Try uploading "
                         "the resume again.")
    # Probe on a dash-normalised copy, never on the block itself. Resumes are
    # full of en dashes (this project's own sample has nine), the model echoes
    # that habit, and a briefing headed "HOW TO SCORE FIT (0-100)" written with
    # U+2013 would otherwise be thrown away as having no rubric at all.
    probe = re.sub(r"[‐-―−]", "-", block).lower()
    if "0-100" not in probe and "score" not in probe:
        raise ValueError("The AI returned a profile with no scoring rubric. "
                         "Try uploading the resume again.")
    block = block[:MAX_BLOCK_CHARS]

    # Resumes are often typeset with the name in caps; the model copies it
    # faithfully, and "SHRUTI GUPTA" then shouts from every screen in the UI.
    name = _clean_str(data.get("name"), 80, "Candidate")
    if name.isupper():
        name = name.title()

    return {
        "name": name,
        "headline": _clean_str(data.get("headline"), 140,
                               name + " - job search profile"),
        "summary": _clean_str(data.get("summary"), 800,
                              "Profile generated from the uploaded resume."),
        "candidate_block": block,
        "target_titles": _clean_list(data.get("target_titles"),
                                     defaults["target_titles"], max_items=24),
        "secondary_titles": _clean_list(data.get("secondary_titles"),
                                        defaults["secondary_titles"],
                                        max_items=24),
        "exclude_title_words": _clean_list(data.get("exclude_title_words"),
                                           defaults["exclude_title_words"],
                                           item_limit=40, max_items=40,
                                           lower=True),
        "skills": _clean_list(data.get("skills"), defaults["skills"],
                              max_items=40),
        "locations": _clean_list(data.get("locations"), defaults["locations"],
                                 max_items=30),
    }


# --------------------------------------------------------------------------
# public entry point
# --------------------------------------------------------------------------

def build_profile(path, today=None):
    """Read a resume and return the profile object. Raises ValueError on failure.

    `today` is an ISO date string; it defaults to the real date at call time
    (never at import time, so a long-running server does not freeze the date on
    the day it booted). The model needs it because the whole rubric hinges on
    whether the newest education end date is in the past or the future.
    """
    if today is None:
        today = datetime.datetime.now().date().isoformat()
    today = str(today).strip() or datetime.datetime.now().date().isoformat()

    resume_text = extract_text(path)
    if len(resume_text) > MAX_RESUME_CHARS:
        resume_text = resume_text[:MAX_RESUME_CHARS] + "\n[...truncated...]"

    # Contact PII leaves this machine only as redaction markers. This is a
    # copy for the prompt alone - the extracted text is never altered.
    prompt_resume = _redact_pii(resume_text)

    key = ai_rank._api_key()
    if not key:
        raise ValueError("No Gemini API key found. Save your key on the setup "
                         "step first, then upload the resume again.")

    # ai_rank._call already asks for responseMimeType application/json and
    # handles 429/5xx backoff. One HTTP client in this project, not two.
    #
    # retries and timeout are pinned instead of inherited. ai_rank's defaults
    # (RETRIES = 5, timeout = 180) are tuned for the batch re-ranker inside a
    # 10-18 minute background run, where 5 x 180s plus ~45s of backoff - about
    # sixteen minutes - is an acceptable worst case. POST /api/resume is
    # synchronous: the browser fetch, and anything proxying it, would abandon
    # the request long before that and leave the user on a spinner that never
    # resolves while the thread kept working on a reply nobody would read.
    # Two attempts at 90s keeps the worst case around three minutes.
    try:
        reply = ai_rank._call(key, _build_prompt(prompt_resume, today),
                              retries=2, timeout=90)
    except Exception as e:
        code = getattr(e, "code", None)
        if code in (400, 401, 403):
            raise ValueError("Gemini rejected the request - the API key looks "
                             "invalid or is not enabled for this model.") from e
        if code == 429:
            raise ValueError("Gemini is rate-limiting this key. Wait a minute, "
                             "then upload the resume again.") from e
        # _call raises RuntimeError("model returned no text (<reason>)") when
        # Gemini answers 200 but returns nothing usable - a safety block, a
        # MAX_TOKENS truncation, or an exhausted quota. That is not a network
        # fault, and "check the connection" sends the user off to fix something
        # that is not broken, so pass the real reason through.
        if isinstance(e, RuntimeError):
            raise ValueError("Gemini did not return a profile for that resume "
                             "(%s). Try uploading it again." % e) from e
        if code:
            raise ValueError("Gemini returned an error (HTTP %s) while reading "
                             "the resume. Try again in a moment." % code) from e
        raise ValueError("Could not reach Gemini to read the resume (%s). Check "
                         "the connection and try again." % type(e).__name__) from e

    return _validate(_parse_object(reply))
